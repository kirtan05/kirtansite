"""Build the /title sample file from one saved survey and the hand-written content.

  cd scripts/title-sample
  uv run --with python-docx --with pillow python build.py --map names.private.json

Writes public/title/{title-file-sample.pdf, title-file-sample-draft.docx, og.png} and
src/data/title-sample.json, then refuses to finish if any real name survives:
every generated artefact (and content.json) is run through the leak detector, and
on any hit the outputs are deleted and the exit code is 1.
"""
import argparse
import base64
import html
import io
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image, ImageFilter

from anon import anonymise_obj, find_leaks, load_map, text_of
from record import load_record

HERE = Path(__file__).resolve().parent
SITE = HERE.parent.parent
CHROME = os.environ.get("CHROME", "/opt/google/chrome/chrome")
FONTS = [("Space Grotesk", "space-grotesk.woff2", "300 700"), ("IBM Plex Mono", "ibm-plex-mono-400.woff2", "400"),
         ("IBM Plex Mono", "ibm-plex-mono-500.woff2", "500"), ("Noto Sans Gujarati", "noto-sans-gujarati-400.woff2", "400"),
         ("Noto Sans Gujarati", "noto-sans-gujarati-600.woff2", "600 700")]
_STATUS = {"પ્રમાણિત": "Certified", "નામંજૂર": "Rejected", "નામંજુર": "Rejected", "રદ": "Cancelled", "પડતર": "Pending"}
e = html.escape
_NO_GLYPH = "\u0ae4"  # e-Dhara types "/" as an unassigned code point that no font draws


def status_en(s: str) -> str:
    return _STATUS.get(s, s)


def area_en(a: str) -> str:
    h, ar, m = (int(x) for x in a.split("-"))
    return f"{a} H-A-Sq.m ({h * 10000 + ar * 100 + m:,} sq m)"


def font_css(font_dir: Path) -> str:
    # Inlined as data URIs: Chrome refuses cross-file font loads from file:// pages.
    out = []
    for family, name, weight in FONTS:
        b64 = base64.b64encode((font_dir / name).read_bytes()).decode()
        out.append(f"@font-face{{font-family:'{family}';src:url(data:font/woff2;base64,{b64}) format('woff2');font-weight:{weight}}}")
    return "\n".join(out)


def blur_thumb(path: str, width: int = 300) -> str:
    im = Image.open(path).convert("L")
    im = im.resize((width, max(1, im.height * width // im.width)))
    im = im.filter(ImageFilter.GaussianBlur(radius=max(2, width // 60)))
    buf = io.BytesIO()
    im.save(buf, "PNG", optimize=True)
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


BASE_CSS = """
@page { size: A4; margin: 16mm 16mm 18mm;
  @bottom-left { content: 'Title File · sample · names, village and survey number changed'; font: 8pt 'IBM Plex Mono', monospace; color: #4B5754; }
  @bottom-right { content: 'kirtanjain.com/title'; font: 8pt 'IBM Plex Mono', monospace; color: #4B5754; } }
* { box-sizing: border-box; }
body { margin: 0; color: #0F1614; font: 10.5pt/1.5 'Space Grotesk', 'Noto Sans Gujarati', sans-serif; }
[lang=gu] { font-family: 'Noto Sans Gujarati', sans-serif; line-height: 1.7; }
.mono { font-family: 'IBM Plex Mono', monospace; }
.page { break-after: page; }
.page:last-child { break-after: auto; }
.wm { position: fixed; top: 42%; left: -10%; right: -10%; text-align: center; transform: rotate(-24deg);
      font: 600 44pt 'Space Grotesk', sans-serif; color: rgba(180, 83, 27, .09); pointer-events: none; }
h1 { font-size: 22pt; margin: 0 0 4mm; letter-spacing: -.01em; }
h2 { font-size: 13pt; margin: 0 0 3mm; padding-bottom: 2mm; border-bottom: 1px solid #C9D3D0; }
h3 { font-size: 10.5pt; margin: 4mm 0 1mm; }
.label { font: 500 8pt 'IBM Plex Mono', monospace; text-transform: uppercase; letter-spacing: .08em; color: #4B5754; }
.tile { border: 1px solid #C9D3D0; border-radius: 3mm; padding: 3mm 4mm; position: relative; margin: 0 0 3mm; }
.tile p { margin: 0 0 1.5mm; }
.tile h3 { margin-top: 0; }
.tile::after { content: ''; position: absolute; inset: 1.3mm; border: 1px dashed #DCE3E1; border-radius: 2mm; }
table { width: 100%; border-collapse: collapse; font-size: 8.5pt; }
th, td { text-align: left; vertical-align: top; padding: 1.1mm 2mm; border-bottom: 1px solid #DCE3E1; }
th { font: 500 7.5pt 'IBM Plex Mono', monospace; text-transform: uppercase; letter-spacing: .06em; color: #4B5754; }
.draft p { font-family: Georgia, 'Times New Roman', serif; font-size: 10.5pt; margin: 0 0 2mm; }
.rejected { color: #B4531B; font-weight: 600; }
.entry { break-inside: avoid; border-bottom: 1px solid #DCE3E1; padding: 2mm 0; }
.scans { display: grid; grid-template-columns: repeat(3, 1fr); gap: 3mm; }
.scans figure { margin: 0; break-inside: avoid; }
.scans img { width: 100%; border: 1px solid #C9D3D0; }
.scans figcaption { font-size: 8pt; margin-top: 1mm; }
.toc li { margin: 1mm 0; }
.note { color: #4B5754; font-size: 9pt; }
"""


def _entry_row(x: dict, content: dict) -> str:
    cls = ' class="rejected"' if status_en(x["status"]) == "Rejected" else ""
    return (f'<tr><td class="mono">{e(x["number"])}</td><td class="mono">{e(x["date"])}</td>'
            f'<td lang="gu">{e(x["type"])}</td><td{cls}>{e(status_en(x["status"]))}</td>'
            f'<td>{e(content["summaries"][x["number"]])}</td></tr>')


def render_sample_html(rec: dict, content: dict, fonts: str, thumbs: dict) -> str:
    where = (f'Survey <span lang="gu">{e(rec["survey"])}</span>, <span lang="gu">{e(rec["village"])}</span>, '
             f'Ta. <span lang="gu">{e(rec["taluka"])}</span>, Dist. <span lang="gu">{e(rec["district"])}</span>')
    points = "".join(f'<div class="tile"><h3>{e(p["title"])}</h3><p>{e(p["detail"])}</p>'
                     f'<p class="label">Entries {e(", ".join(p["entries"]))}</p></div>' for p in content["points"])
    draft = "".join(f'<h3>{e(d["item"])}. {e(d["heading"])}</h3>' + "".join(f"<p>{e(par)}</p>" for par in d["text"].split("\n\n"))
                    for d in content["draft"])
    owners = "".join(f'<li><span lang="gu">{e(o["name"])}</span> <span class="mono">(entry {e(o["entry"])})</span></li>' for o in rec["owners"])
    boja = "".join(f'<li lang="gu">{e(n)}</li>' for n in rec["boja_notes"]) or "<li>None recorded</li>"
    rows = "".join(_entry_row(x, content) for x in rec["entries"])
    full = "".join(f'<div class="entry"><p class="label">Entry {e(x["number"])} · {e(x["date"])} · {e(status_en(x["status"]))}</p>'
                   f'<p lang="gu"><b>{e(x["type"])}</b> — {e(x["text"].replace(_NO_GLYPH, "/"))}</p><p lang="gu" class="note">{e(x["remark"])}</p></div>'
                   for x in rec["entries"])
    scans = "".join(f'<figure><img src="{src}" alt="Handwritten entry {e(o["number"])}, blurred">'
                    f'<figcaption><span class="mono">Entry {e(o["number"])}</span>: {e(content["old_summaries"][o["number"]])}</figcaption></figure>'
                    for o in rec["old_entries"] for src in thumbs[o["number"]][:1])
    deeds = "".join(f'<tr><td>{e(d["sro"])}</td><td class="mono">••••/{e(d["year"])}</td><td class="mono">{e(d["date"])}</td>'
                    f'<td lang="gu">{"<br>".join(e(p["role"] + ": " + p["name"]) for p in d["parties"])}</td></tr>' for d in rec["deeds"])
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Title File sample</title>
<style>{fonts}{BASE_CSS}</style></head><body>
<div class="wm" aria-hidden="true">SAMPLE: names changed</div>

<section class="page">
<p class="label">Title File · <span lang="gu">ટાઇટલ ફાઇલ</span></p>
<h1>Revenue file and draft</h1>
<p>{where}</p>
<p>Records as of <span class="mono">{e(rec["as_of"])}</span> · Bank format: Bank of Baroda, Title Opinion Report</p>
<div class="tile"><p class="label">Contents</p><ol class="toc">
<li>Summary and points to check</li><li>Draft: revenue-record items 1, 2, 5–9, 14 and 15</li>
<li>Annex A: current record</li><li>Annex B: Village Form 6 entries, with English summaries</li>
<li>Annex C: full text of typed entries</li><li>Annex D: handwritten entries (blurred in this sample)</li>
<li>Annex E: registered documents</li></ol></div>
<p class="note">Real record. Names, village and survey number changed. A records and drafting service, not legal advice.</p>
</section>

<section class="page"><h2>Summary and points to check</h2>
<p>{len(rec["entries"])} typed entries from {e(rec["entries"][0]["date"])} to {e(rec["entries"][-1]["date"])}, and {len(rec["old_entries"])} handwritten entries.
Current holder{"s" if len(rec["owners"]) > 1 else ""}: {", ".join(f'<span lang="gu">{e(o["name"])}</span>' for o in rec["owners"])}.</p>
{points}</section>

<section class="page draft"><h2>Draft for the advocate's review</h2>
<p class="note">Revenue-record items of the Bank of Baroda Title Opinion Report. Items 3, 4, 10–13 and 16–20 need the originals, the sub-registrar search and your opinion.</p>
{draft}</section>

<section class="page"><h2>Annex A: current record</h2><table>
<tr><th>District</th><td lang="gu">{e(rec["district"])}</td></tr><tr><th>Taluka</th><td lang="gu">{e(rec["taluka"])}</td></tr>
<tr><th>Village</th><td lang="gu">{e(rec["village"])}</td></tr><tr><th>Survey</th><td lang="gu">{e(rec["survey"])}</td></tr>
<tr><th>UPIN</th><td class="mono">{e(rec["upin"])}</td></tr><tr><th>Khata</th><td class="mono">{e(rec["khata"])}</td></tr>
<tr><th>Area</th><td class="mono">{e(area_en(rec["area"]))}</td></tr><tr><th>Assessment</th><td class="mono">Rs. {e(rec["assessment"])}</td></tr>
<tr><th>Tenure</th><td lang="gu">{e(rec["tenure"])}</td></tr><tr><th>Land use</th><td lang="gu">{e(rec["land_use"])}</td></tr>
<tr><th>Holders</th><td><ul>{owners}</ul></td></tr><tr><th>Charges and other rights</th><td><ul>{boja}</ul>
<p class="mono">Entries {e(", ".join(rec["boja_entries"]))}</p></td></tr>
<tr><th>Tenants</th><td>{e(", ".join(rec["tenants"])) or "None recorded"}</td></tr>
<tr><th>Form 8-A</th><td>Included in real files</td></tr></table></section>

<section class="page"><h2>Annex B: Village Form 6 entries</h2><table>
<tr><th>Entry</th><th>Date</th><th>Type</th><th>Status</th><th>Summary</th></tr>{rows}</table></section>

<section class="page"><h2>Annex C: full text of typed entries</h2>{full}</section>

<section class="page"><h2>Annex D: handwritten entries</h2>
<p class="note">Real files include the full scans. They are blurred here because they show real names.</p>
<div class="scans">{scans}</div></section>

<section class="page"><h2>Annex E: registered documents</h2><table>
<tr><th>Office</th><th>Doc. no.</th><th>Date</th><th>Parties</th></tr>{deeds}</table>
<p class="note">Registered transactions from 2007 onward (Index-2): included in real files. Document numbers are masked in this sample.</p>
</section></body></html>"""


def render_og_html(rec: dict, content: dict, fonts: str) -> str:
    quote = content["chain_excerpt"][:230].rsplit(" ", 1)[0] + " …"
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><style>{fonts}
* {{ box-sizing: border-box; }}
body {{ margin: 0; width: 1200px; height: 630px; background: #E9EEEC; color: #0F1614; font-family: 'Space Grotesk', sans-serif;
       display: grid; grid-template-columns: 1fr 1fr; gap: 40px; padding: 64px; }}
[lang=gu] {{ font-family: 'Noto Sans Gujarati', sans-serif; }}
.label {{ font: 500 20px 'IBM Plex Mono', monospace; letter-spacing: .08em; text-transform: uppercase; color: #4B5754; }}
h1 {{ font-size: 54px; line-height: 1.08; margin: 18px 0 20px; letter-spacing: -.02em; }}
.gu {{ font-size: 26px; line-height: 1.5; color: #4B5754; }}
.price {{ margin-top: 28px; font: 500 28px 'IBM Plex Mono', monospace; color: #B4531B; }}
.tile {{ background: #FBFCFB; border: 2px solid #C9D3D0; border-radius: 16px; padding: 30px; position: relative; align-self: center; }}
.tile::after {{ content: ''; position: absolute; inset: 7px; border: 2px dashed #DCE3E1; border-radius: 11px; }}
.tile p {{ font: 22px/1.5 Georgia, serif; margin: 12px 0 0; }}
</style></head><body>
<div><div class="label">Title File · <span lang="gu">ટાઇટલ ફાઇલ</span></div>
<h1>30-year revenue file and draft chain of title</h1>
<div class="gu" lang="gu">સરવે નંબર આપો, બીજા દિવસે સાંજ સુધીમાં ફાઇલ.</div>
<div class="price">₹199 · first 3 free</div></div>
<div class="tile"><div class="label">Item 15 · chain of title</div><p>{e(quote)}</p></div>
</body></html>"""


def excerpt(rec: dict, content: dict, pdf_pages: int) -> dict:
    by_no = {x["number"]: x for x in rec["entries"]}
    rows = [{"number": n, "date": by_no[n]["date"], "type": by_no[n]["type"], "status": by_no[n]["status"],
             "status_en": status_en(by_no[n]["status"]), "summary": content["summaries"][n]} for n in content["excerpt_rows"]]
    points = [{"title": content["points"][i]["title"], "detail": content["points"][i]["detail"]} for i in content["excerpt_points"]]
    return {"survey": rec["survey"], "village": rec["village"], "taluka": rec["taluka"], "district": rec["district"],
            "as_of": rec["as_of"], "pdf_pages": pdf_pages, "chain_excerpt": content["chain_excerpt"], "rows": rows, "points": points}


def write_docx(path: Path, rec: dict, content: dict) -> None:
    from docx import Document  # only the full build needs python-docx
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(11)
    doc.sections[0].header.paragraphs[0].text = "SAMPLE: names changed. Draft for the advocate's review."
    doc.add_heading("Revenue-record items: Bank of Baroda Title Opinion Report", level=1)
    doc.add_paragraph(f"Survey {rec['survey']}, {rec['village']}, Ta. {rec['taluka']}, Dist. {rec['district']}. "
                      f"Revenue records as of {rec['as_of']}.")
    for d in content["draft"]:
        doc.add_heading(f"{d['item']}. {d['heading']}", level=2)
        for par in d["text"].split("\n\n"):
            doc.add_paragraph(par)
    doc.save(path)


def _chrome(*args: str) -> None:
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-sandbox", "--hide-scrollbars",
                    "--virtual-time-budget=4000", *args], check=True, capture_output=True)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--map", default=str(HERE / "names.private.json"))
    args = ap.parse_args()
    m = load_map(args.map)
    content = json.loads((HERE / "content.json").read_text(encoding="utf-8"))
    raw = load_record(Path(m["source"]))
    rec = anonymise_obj(raw, m)
    fonts = font_css(SITE / "public/fonts/title")
    thumbs = {o["number"]: [blur_thumb(s) for s in o["scans"]] for o in raw["old_entries"]}

    build, out = HERE / "build", SITE / "public/title"
    build.mkdir(exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    (build / "sample.html").write_text(render_sample_html(rec, content, fonts, thumbs), encoding="utf-8")
    (build / "og.html").write_text(render_og_html(rec, content, fonts), encoding="utf-8")

    pdf, png, docx = out / "title-file-sample.pdf", out / "og.png", out / "title-file-sample-draft.docx"
    _chrome("--no-pdf-header-footer", f"--print-to-pdf={pdf}", (build / "sample.html").as_uri())
    _chrome("--window-size=1200,630", f"--screenshot={png}", (build / "og.html").as_uri())
    write_docx(docx, rec, content)
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True, check=True).stdout
    pages = int(re.search(r"Pages:\s+(\d+)", info).group(1))
    data = SITE / "src/data/title-sample.json"
    data.write_text(json.dumps(excerpt(rec, content, pages), ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    # Leak gate: every artefact that can reach the public repo. The PDF is checked through its extracted text.
    pdf_text = subprocess.run(["pdftotext", str(pdf), "-"], capture_output=True, text=True, check=True).stdout
    checked = {f.name: text_of(f) for f in [HERE / "content.json", build / "sample.html", build / "og.html", docx, data]}
    checked[pdf.name] = pdf_text
    leaks = {k: find_leaks(v, m) for k, v in checked.items()}
    leaks = {k: v for k, v in leaks.items() if v}
    if leaks:
        for f in (pdf, png, docx, data):
            f.unlink(missing_ok=True)
        print("LEAKS, outputs deleted:", json.dumps(leaks, ensure_ascii=False), file=sys.stderr)
        return 1
    print(f"no leaks in {len(checked)} artefacts")
    print(f"ok: {pdf.name} ({pages} pages, {pdf.stat().st_size // 1024} KB), {png.name} ({png.stat().st_size // 1024} KB), {docx.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
